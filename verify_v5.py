"""
Rigorous Hostile-Judge Verification Script for SupplyGraph_Final_v5.pptx and SupplyGraph_Blue_Final.pptx
"""

import sys
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt

BANNED_STRINGS = [
    "0x5FbDB",
    "Fabric",
    "Zero-FP",
    "tree-sitter",
    "1.5",
    "zero-host",
    "detonat",
    "Jenkins",
    "...",
    "TODO",
    "[cite",
]

def verify_deck(file_path):
    print(f"\n=======================================================")
    print(f"AUDITING: {os.path.basename(file_path)}")
    print(f"=======================================================")

    prs = Presentation(file_path)
    slides = list(prs.slides)
    errors = []
    warnings = []

    # 1. Slide Count
    if len(slides) != 6:
        errors.append(f"Expected exactly 6 slides, got {len(slides)}")
    else:
        print(f" [PASS] Exact slide count: {len(slides)} slides")

    # 2. Dimensions
    w_in = prs.slide_width.inches
    h_in = prs.slide_height.inches
    if abs(w_in - 13.333) > 0.05 or abs(h_in - 7.5) > 0.05:
        errors.append(f"Invalid dimensions: {w_in:.2f}x{h_in:.2f} (expected 13.33x7.50)")
    else:
        print(f" [PASS] Dimensions: {w_in:.2f} x {h_in:.2f} inches (16:9 widescreen)")

    # 3. Check for Banned Strings & Brackets
    all_text = []
    for i, slide in enumerate(slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        slide_text.append(t)
                        # Check banned strings
                        for banned in BANNED_STRINGS:
                            if banned.lower() in t.lower():
                                errors.append(f"Slide {i}: Banned string '{banned}' found in: '{t}'")
                        # Check for square brackets in matter text
                        if "[" in t or "]" in t:
                            errors.append(f"Slide {i}: Square bracket '[' or ']' found in: '{t}'")

        full_s_text = " ".join(slide_text)
        all_text.append(full_s_text)

    deck_str = " ".join(all_text)

    # 4. Required Verified Claims
    required_checks = [
        ("Gemini 2.5 Flash", "Gemini 2.5 Flash pinned"),
        ("130/130", "130/130 passing tests"),
        ("74%", "74% full-suite line coverage"),
        ("Token-Sanitized", "Token-Sanitized Pattern Engine for JS/TS"),
        ("Python AST", "Python AST stdlib parser"),
        ("Tarjan SCC", "Tarjan SCC condensation"),
        ("min(100, 2P+15S+25A+50D)", "Blast Radius formula"),
        ("1-Prod(1-wi*ri)", "Confidence formula"),
        ("3380e39", "Git commit hash 3380e39"),
        ("docs/EVIDENCE.md", "Link to docs/EVIDENCE.md"),
        ("Target: Polygon Amoy Testnet", "Honest Amoy status"),
    ]

    for req, desc in required_checks:
        if req.lower() in deck_str.lower():
            print(f" [PASS] Verified claim present: {desc} ('{req}')")
        else:
            errors.append(f"Missing required claim: {desc} ('{req}')")

    # 5. Slide 4 Competitor Matrix clearance
    slide4 = slides[3]
    for shape in slide4.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Snyk: call-graph" in p.text:
                    bottom_y = shape.top + shape.height
                    clearance = (Inches(7.5 - 0.28) - bottom_y) / 914400.0
                    if clearance < 0.10:
                        warnings.append(f"Slide 4 footnote clearance is only {clearance:.2f} in (<0.10 in)")
                    else:
                        print(f" [PASS] Slide 4 footnote clearance above footer: {clearance:.2f} in (>=0.10 in)")

    print(f"\n--- AUDIT SUMMARY FOR {os.path.basename(file_path)} ---")
    if errors:
        print(f"[FAIL] FAILED with {len(errors)} error(s):")
        for e in errors:
            print(f"   - {e}")
        return False
    else:
        print(f"[SUCCESS] PASSED with 0 ERRORS! (Warnings: {len(warnings)})")
        for w in warnings:
            print(f"   [WARN] {w}")
        return True


if __name__ == "__main__":
    v5_path = r"c:\Users\DARAM  HEMANTH REDDY\Desktop\supplygraph\SupplyGraph_Final_v5.pptx"
    blue_path = r"c:\Users\DARAM  HEMANTH REDDY\Desktop\supplygraph\SupplyGraph_Blue_Final.pptx"

    ok1 = verify_deck(v5_path)
    ok2 = verify_deck(blue_path)

    if ok1 and ok2:
        print("\nALL PRESENTATIONS AUDITED AND CERTIFIED 100% CLEAN!")
        sys.exit(0)
    else:
        sys.exit(1)
